"""Generator for table (ListObject/structured reference) test cases (Tier 3)."""

from pathlib import Path
from typing import Any

import xlwings as xw

from excelbench.generator.base import FeatureGenerator
from excelbench.models import Importance, TestCase


class TablesGenerator(FeatureGenerator):
    """Generates test cases for Excel tables (ListObjects)."""

    feature_name = "tables"
    tier = 3
    filename = "19_tables.xlsx"

    def __init__(self) -> None:
        self._ops: list[dict[str, object]] = []

    def generate(self, sheet: xw.Sheet) -> list[TestCase]:
        self.setup_header(sheet)

        test_cases: list[TestCase] = []
        row = 2

        def _add_case(
            *,
            case_id: str,
            label: str,
            name: str,
            ref: str,
            columns: list[str],
            style: str | None,
            totals_row: bool = False,
            autofilter: bool = False,
            importance: Importance = Importance.BASIC,
        ) -> None:
            expected: dict[str, Any] = {
                "table": {
                    "name": name,
                    "ref": ref,
                    "header_row": True,
                    "totals_row": totals_row,
                    "style": style,
                    "columns": columns,
                }
            }
            expected_table = expected["table"]
            if totals_row:
                expected_table["totals_row_count"] = 1
            if autofilter:
                expected_table["autofilter"] = True

            self.write_test_case(sheet, row, label, expected)
            test_cases.append(
                TestCase(
                    id=case_id,
                    label=label,
                    row=row,
                    expected=expected,
                    importance=importance,
                )
            )
            self._ops.append(
                {
                    "name": name,
                    "ref": ref,
                    "style": style,
                    "columns": columns,
                    "totals_row": totals_row,
                    "autofilter": autofilter,
                }
            )

        # Keep table bodies away from A:C. Column C stores JSON expectations.

        # 1) Basic 3-column table
        label = "Table: basic 3-col"
        sheet.range("E2").value = "Name"
        sheet.range("F2").value = "Qty"
        sheet.range("G2").value = "Price"
        sheet.range("E3").value = "Widget"
        sheet.range("F3").value = 10
        sheet.range("G3").value = 4.99
        sheet.range("E4").value = "Gadget"
        sheet.range("F4").value = 5
        sheet.range("G4").value = 12.50
        sheet.range("E5").value = "Gizmo"
        sheet.range("F5").value = 8
        sheet.range("G5").value = 7.25
        _add_case(
            case_id="tbl_basic",
            label=label,
            name="SalesData",
            ref="E2:G5",
            columns=["Name", "Qty", "Price"],
            style="TableStyleMedium9",
        )
        row += 1

        # 2) Table with totals row
        label = "Table: with totals row"
        sheet.range("E7").value = "Item"
        sheet.range("F7").value = "Count"
        sheet.range("G7").value = "Total"
        sheet.range("E8").value = "Apples"
        sheet.range("F8").value = 2
        sheet.range("G8").value = 4.00
        sheet.range("E9").value = "Oranges"
        sheet.range("F9").value = 3
        sheet.range("G9").value = 7.50
        sheet.range("E10").value = "Bananas"
        sheet.range("F10").value = 1
        sheet.range("G10").value = 1.25
        sheet.range("E11").value = "Total"
        _add_case(
            case_id="tbl_with_totals",
            label=label,
            name="Summary",
            ref="E7:G11",
            columns=["Item", "Count", "Total"],
            style="TableStyleLight1",
            totals_row=True,
        )
        row += 1

        # 3) Table with no style
        label = "Table: no style"
        sheet.range("E13").value = "Key"
        sheet.range("F13").value = "Value"
        sheet.range("E14").value = "a"
        sheet.range("F14").value = 1
        sheet.range("E15").value = "b"
        sheet.range("F15").value = 2
        sheet.range("E16").value = "c"
        sheet.range("F16").value = 3
        _add_case(
            case_id="tbl_no_style",
            label=label,
            name="PlainTable",
            ref="E13:F16",
            columns=["Key", "Value"],
            style=None,
        )
        row += 1

        # 4) Single-column table
        label = "Table: single column"
        sheet.range("E18").value = "Score"
        sheet.range("E19").value = 10
        sheet.range("E20").value = 20
        sheet.range("E21").value = 30
        _add_case(
            case_id="tbl_single_col",
            label=label,
            name="SingleCol",
            ref="E18:E21",
            columns=["Score"],
            style="TableStyleMedium2",
            importance=Importance.EDGE,
        )
        row += 1

        # 5) Header-only table (no data rows)
        label = "Table: header only (no data rows)"
        sheet.range("E23").value = "A"
        sheet.range("F23").value = "B"
        sheet.range("G23").value = "C"
        _add_case(
            case_id="tbl_single_row",
            label=label,
            name="EmptyTable",
            ref="E23:G23",
            columns=["A", "B", "C"],
            style="TableStyleMedium9",
            importance=Importance.EDGE,
        )
        row += 1

        # 6) Table with autoFilter
        label = "Table: with autoFilter"
        sheet.range("E25").value = "Region"
        sheet.range("F25").value = "Sales"
        sheet.range("G25").value = "Year"
        sheet.range("E26").value = "NA"
        sheet.range("F26").value = 100
        sheet.range("G26").value = 2024
        sheet.range("E27").value = "EU"
        sheet.range("F27").value = 200
        sheet.range("G27").value = 2025
        sheet.range("E28").value = "APAC"
        sheet.range("F28").value = 150
        sheet.range("G28").value = 2026
        _add_case(
            case_id="tbl_autofilter",
            label=label,
            name="Filtered",
            ref="E25:G28",
            columns=["Region", "Sales", "Year"],
            style="TableStyleMedium9",
            autofilter=True,
            importance=Importance.EDGE,
        )

        return test_cases

    def post_process(self, output_path: Path) -> None:
        # xlwings does not create real ListObject definitions reliably across
        # platforms. Post-save using openpyxl for reproducibility.
        if not self._ops:
            return

        from openpyxl import load_workbook
        from openpyxl.worksheet.filters import AutoFilter
        from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

        wb = load_workbook(output_path)
        try:
            ws = wb[self.feature_name]
            for op in self._ops:
                name = str(op["name"])
                ref = str(op["ref"])
                style_name = op.get("style")
                totals = bool(op.get("totals_row", False))
                autofilter = bool(op.get("autofilter", False))
                columns = op.get("columns")

                style = None
                if style_name:
                    style = TableStyleInfo(
                        name=str(style_name),
                        showFirstColumn=False,
                        showLastColumn=False,
                        showRowStripes=True,
                        showColumnStripes=False,
                    )

                table = Table(displayName=name, ref=ref)
                if style is not None:
                    table.tableStyleInfo = style
                if totals:
                    table.totalsRowCount = 1
                if isinstance(columns, list) and columns:
                    table.tableColumns = [
                        TableColumn(id=i + 1, name=str(col)) for i, col in enumerate(columns)
                    ]
                if autofilter:
                    table.autoFilter = AutoFilter(ref=ref)

                ws.add_table(table)

            wb.save(output_path)
        finally:
            wb.close()
